from docx import Document
p=r'C:\Users\anh15\Downloads\SD-41_DATNSP26_Hoan_Thanh.docx'
d=Document(p)
print('PARAS',len(d.paragraphs),'TABLES',len(d.tables),'SECTIONS',len(d.sections))
for i,p in enumerate(d.paragraphs):
 t=p.text.strip()
 if t: print(f'{i}: {t[:240]}')
