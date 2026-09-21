import re,json,collections
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[3]
OUT=ROOT/'output/report/schema-43'
OUT.mkdir(parents=True,exist_ok=True)
raw=(ROOT/'tmp/pdfs/report-audit/schema.sql').read_text(encoding='utf-8-sig')
def split(s):
    result=[];start=0;depth=0;quote=False
    for i,c in enumerate(s):
        if c=="'": quote=not quote
        if not quote:
            if c=='(':depth+=1
            elif c==')':depth-=1
            elif c==',' and depth==0:result.append(s[start:i].strip());start=i+1
    return result+[s[start:].strip()]
tables={};fks=[]
for line in raw.splitlines():
    m=re.search(r'create table (\w+) \((.*)\)$',line)
    if m:
        name,body=m.groups();cols=[];pk=[];unique=[]
        for part in split(body):
            if part.startswith('primary key'):pk=re.findall(r'\w+',part[part.index('(')+1:-1]);continue
            if part.startswith('unique'):unique.append([x.strip() for x in part[part.index('(')+1:-1].split(',')]);continue
            c,rest=part.split(' ',1)
            mt=re.match(r'(enum \(.*?\)|\w+(?:\(\d+\))?)(.*)',rest)
            typ,extra=mt.groups()
            cols.append(dict(name=c,type=typ,nullable='not null' not in extra,identity='identity' in extra,unique='unique' in extra,extra=extra.strip()))
            if 'unique' in extra:unique.append([c])
        for col in cols:
            col['pk']=col['name'] in pk
            if col['pk']:col['nullable']=False
        tables[name]=dict(columns=cols,pk=pk,unique=unique)
    m=re.search(r'alter table if exists (\w+) add constraint \w+ foreign key \((\w+)\) references (\w+)',line)
    if m:fks.append(m.groups())
assert len(tables)==43 and len(fks)==39,(len(tables),len(fks))

purposes={
'roles':'Vai trò phân quyền','users':'Tài khoản người dùng','user_profiles':'Hồ sơ thể lực và điều kiện tập luyện','memberships':'Gói hội viên đã đăng ký','invoices':'Hóa đơn gói hội viên và vật phẩm trang trí',
'exercises':'Thư viện bài tập','foods':'Thư viện món ăn','endurance_tests':'Kết quả kiểm tra sức bền hiện tại của từng người dùng','progress_tracking':'Lịch sử chỉ số cơ thể','workout_plans':'Giáo án cá nhân và giáo án mẫu','workout_plan_days':'Ngày tập trong giáo án','workout_plan_exercises':'Bài tập và định mức trong ngày tập','plan_muscle_group_weight':'Hệ số ưu tiên nhóm cơ theo giáo án','workout_sessions':'Buổi tập thực tế','session_exercise_logs':'Kết quả từng bài tập trong buổi','weekly_reviews':'Đánh giá giáo án theo tuần',
'muscle_split_configs':'Cấu hình chia nhóm cơ theo mục tiêu/số buổi','recommended_schedule_configs':'Cấu hình ngày tập đề xuất theo số buổi','injury_area_options':'Danh mục vùng chấn thương','system_configs':'Cấu hình số của hệ thống',
'chat_messages':'Lịch sử hội thoại người dùng với chatbot','notifications':'Thông báo người dùng','service_ratings':'Đánh giá dịch vụ và phản hồi quản trị','support_sessions':'Phiên hỗ trợ trực tiếp','support_messages':'Tin nhắn trong phiên hỗ trợ','pet_profiles':'Trạng thái nhân vật đồng hành','user_cosmetic_ownership':'Vật phẩm trang trí người dùng sở hữu','work_shifts':'Ca làm và đối soát tiền của nhân viên',
'shop_products':'Sản phẩm cửa hàng','product_attributes':'Danh mục thuộc tính sản phẩm','product_attribute_values':'Giá trị của thuộc tính','product_variants':'Biến thể sản phẩm','variant_attribute_values':'Bảng nối biến thể với các giá trị thuộc tính','shop_cart_items':'Dòng giỏ hàng của tài khoản','shop_orders':'Đơn hàng online và tại quầy','shop_order_items':'Chi tiết hàng hóa được chốt trong đơn','vouchers':'Mã giảm giá và điều kiện áp dụng','voucher_scope_products':'Danh sách ID sản phẩm thuộc phạm vi voucher','product_reviews':'Đánh giá sản phẩm theo đơn mua','customer_addresses':'Sổ địa chỉ nhận hàng của khách','customer_product_states':'Yêu thích và lần xem gần nhất trên cùng một dòng','shop_order_events':'Lịch sử trạng thái và sự kiện đơn hàng','shop_inventory_movements':'Nhật ký thay đổi tồn kho'}
assert set(purposes)==set(tables)
groups={
'Tài khoản và hội viên':['roles','users','user_profiles','memberships','invoices'],
'Tập luyện và dinh dưỡng':['exercises','foods','endurance_tests','progress_tracking','workout_plans','workout_plan_days','workout_plan_exercises','plan_muscle_group_weight','workout_sessions','session_exercise_logs','weekly_reviews'],
'Cấu hình':['muscle_split_configs','recommended_schedule_configs','injury_area_options','system_configs'],
'Tương tác và hỗ trợ':['chat_messages','notifications','service_ratings','support_sessions','support_messages'],
'Nhân vật và nhân viên':['pet_profiles','user_cosmetic_ownership','work_shifts'],
'Cửa hàng':['shop_products','product_attributes','product_attribute_values','product_variants','variant_attribute_values','shop_cart_items','shop_orders','shop_order_items','vouchers','voucher_scope_products','product_reviews','customer_addresses','customer_product_states','shop_order_events','shop_inventory_movements']}
logical=[('customer_addresses','user_id','users'),('customer_product_states','user_id','users'),('customer_product_states','product_id','shop_products'),('shop_order_events','order_id','shop_orders'),('shop_inventory_movements','product_id','shop_products'),('shop_inventory_movements','variant_id','product_variants'),('shop_inventory_movements','order_id','shop_orders'),('shop_cart_items','variant_id','product_variants'),('shop_order_items','product_id','shop_products'),('shop_order_items','variant_id','product_variants'),('shop_orders','created_by_staff_id','users'),('voucher_scope_products','product_id','shop_products'),('user_cosmetic_ownership','user_id','users'),('workout_plans','original_plan_id','workout_plans'),('workout_plan_exercises','last_low_adjustment_log_id','session_exercise_logs')]
for t,c,p in logical:assert c in [x['name'] for x in tables[t]['columns']]

# Read entity sources to retain field comments and prove every physical table has a mapping.
sources={};comments={}
base=ROOT/'gym-management/src/main/java'
def snake(s):return re.sub(r'(?<!^)(?=[A-Z])','_',s).lower()
for p in base.rglob('*.java'):
    s=p.read_text(encoding='utf-8-sig')
    for m in re.finditer(r'@(?:Table|JoinTable|CollectionTable)\s*\(\s*name\s*=\s*"([^"]+)"',s):
        if m[1] in tables:sources[m[1]]=str(p.relative_to(ROOT)).replace('\\','/')
    tm=re.search(r'@Table\s*\(\s*name\s*=\s*"([^"]+)"',s)
    if tm:
        for m in re.finditer(r'private\s+[\w<>?, ]+\s+(\w+)\s*(?:=[^;]*)?;\s*//([^\n]*)',s):comments[(tm[1],snake(m[1]))]=m[2].strip()
assert set(sources)==set(tables),set(tables)-set(sources)

labels={
'id':'Mã định danh','name':'Tên hiển thị','code':'Mã nghiệp vụ','description':'Nội dung mô tả','notes':'Ghi chú','note':'Ghi chú','status':'Trạng thái xử lý','created_at':'Thời điểm tạo','updated_at':'Thời điểm cập nhật','user_id':'Mã người dùng','role_id':'Mã vai trò','email':'Địa chỉ email','password':'Mật khẩu đã mã hóa','full_name':'Họ và tên','phone':'Số điện thoại','role_name':'Tên vai trò','email_verified':'Đã xác minh email','verification_token':'Mã xác minh tài khoản','password_reset_attempts':'Số lần thử đặt lại mật khẩu','password_reset_blocked_until':'Thời hạn chặn đặt lại mật khẩu',
'product_id':'Mã sản phẩm','variant_id':'Mã biến thể','order_id':'Mã đơn hàng','attribute_id':'Mã thuộc tính','value_id':'Mã giá trị thuộc tính','voucher_id':'Mã voucher','membership_id':'Mã đăng ký hội viên','workout_plan_id':'Mã giáo án','plan_day_id':'Mã ngày tập','session_id':'Mã buổi tập hoặc phiên hỗ trợ','exercise_id':'Mã bài tập','admin_id':'Mã người hỗ trợ','created_by_staff_id':'Mã nhân viên tạo đơn','original_plan_id':'Mã giáo án nguồn','last_low_adjustment_log_id':'Mã log đã dùng cho lần điều chỉnh giảm',
'price':'Giá tiền','sale_price':'Giá bán ưu đãi','price_override':'Giá riêng của biến thể','stock':'Số lượng tồn','active':'Đang hoạt động','is_active':'Đang hoạt động','sku':'Mã SKU biến thể','image_url':'Đường dẫn ảnh đại diện','images':'Danh sách URL ảnh, JSON trong VARCHAR(20000)','category':'Nhóm phân loại','brand':'Thương hiệu','suitable_goals':'Các mục tiêu phù hợp','required_equipment_code':'Mã thiết bị liên quan','quantity':'Số lượng','unit_price':'Đơn giá tại thời điểm mua','line_total':'Thành tiền của dòng','product_name':'Tên sản phẩm tại thời điểm mua','variant_label':'Tên phân loại tại thời điểm mua/thêm giỏ','subtotal':'Tổng tiền hàng','discount':'Khoản giảm giá','shipping_fee':'Phí vận chuyển','total':'Tổng phải trả','voucher_discount':'Số tiền được giảm bằng voucher','voucher_code':'Mã voucher chốt trên đơn','channel':'Kênh bán','payment_method':'Phương thức thanh toán','payment_url':'URL thanh toán của cổng','gateway_reference':'Mã tham chiếu cổng thanh toán','paid_at':'Thời điểm xác nhận đã trả tiền','expires_at':'Thời điểm hết hạn','cancelled_at':'Thời điểm hủy','transaction_id':'Mã giao dịch','transfer_code':'Mã nội dung chuyển khoản','qr_code_url':'URL ảnh QR','qr_raw_payload':'Dữ liệu QR gốc','receiver_name':'Tên người nhận','shipping_address':'Địa chỉ giao chốt trên đơn','address':'Địa chỉ nhận hàng','default_address':'Địa chỉ mặc định','wishlisted':'Được đánh dấu yêu thích','viewed_at':'Lần xem gần nhất','actor':'Người/nguồn thực hiện','reason':'Lý do thay đổi','before_stock':'Tồn trước thay đổi','after_stock':'Tồn sau thay đổi','from_status':'Trạng thái trước sự kiện','to_status':'Trạng thái sau sự kiện','rating':'Điểm đánh giá','comment':'Nội dung đánh giá','value':'Giá trị thuộc tính hoặc mức giảm','type':'Loại nghiệp vụ','scope_type':'Loại phạm vi áp dụng','scope_category':'Danh mục áp dụng','min_order_amount':'Giá trị đơn tối thiểu','max_discount_amount':'Mức giảm tối đa','usage_limit':'Giới hạn lượt sử dụng','used_count':'Số lượt đã dùng','start_at':'Thời điểm bắt đầu hiệu lực','end_at':'Thời điểm kết thúc hiệu lực',
'height':'Chiều cao','weight':'Cân nặng','initial_weight':'Cân nặng ban đầu','bmi':'Chỉ số BMI','body_fat_percentage':'Tỷ lệ mỡ cơ thể','muscle_mass_kg':'Khối lượng cơ (kg)','arm_cm':'Vòng tay (cm)','chest_cm':'Vòng ngực (cm)','hip_cm':'Vòng hông (cm)','thigh_cm':'Vòng đùi (cm)','waist_cm':'Vòng eo (cm)','recorded_at':'Thời điểm ghi nhận','recorded_date':'Ngày ghi nhận','source':'Nguồn ghi nhận','age':'Tuổi','gender':'Giới tính','date_of_birth':'Ngày sinh','goal':'Mục tiêu tập luyện','fitness_level':'Mức thể lực','training_location':'Địa điểm tập','daily_activity_level':'Mức vận động hằng ngày','available_equipment':'Danh sách thiết bị hiện có','disliked_exercises':'Danh sách bài tập không thích','injury_areas':'Các vùng chấn thương','medical_conditions':'Thông tin sức khỏe','available_days_per_week':'Số ngày có thể tập mỗi tuần','preferred_session_duration':'Thời lượng buổi tập mong muốn','preferred_training_days':'Các ngày tập mong muốn','training_experience_months':'Kinh nghiệm tập (tháng)',
'plank_seconds':'Thành tích plank (giây)','pushup_reps':'Số lần chống đẩy','squat_reps':'Số lần squat','tested_at':'Thời điểm kiểm tra','calories':'Năng lượng món ăn','protein_grams':'Lượng protein (g)','fat_grams':'Lượng chất béo (g)','weight_grams':'Khối lượng khẩu phần (g)','ingredients':'Nguyên liệu','instructions':'Hướng dẫn chế biến','multiplier':'Hệ số ưu tiên','muscle_group':'Nhóm cơ','secondary_muscle_groups':'Nhóm cơ phụ','contraindicated_injuries':'Các chấn thương chống chỉ định','video_url':'URL video bài tập','difficulty':'Độ khó','assessment_metric_type':'Chỉ tiêu kiểm tra','uses_weight':'Có sử dụng tạ','is_assessment':'Là bài kiểm tra','stamina_cost':'Chi phí sức bền','rest_seconds':'Thời gian nghỉ (giây)',
'sessions_per_week':'Số buổi mỗi tuần','recommended_days':'Các ngày tập đề xuất','day_groups':'Cấu hình nhóm cơ cho các ngày','label':'Nhãn hiển thị','config_key':'Khóa cấu hình','config_value':'Giá trị số cấu hình','content':'Nội dung tin nhắn','sender':'Nguồn gửi USER/BOT','sender_role':'Vai trò bên gửi','title':'Tiêu đề','message':'Nội dung thông báo','is_read':'Đã đọc','ref_id':'Mã đối tượng tham chiếu theo ref_type','ref_type':'Loại đối tượng tham chiếu','scheduled_at':'Thời điểm dự kiến gửi','sent_at':'Thời điểm đã gửi','subject':'Chủ đề hỗ trợ','accepted_at':'Thời điểm tiếp nhận','closed_at':'Thời điểm đóng','rated_at':'Thời điểm đánh giá','user_rating':'Điểm khách đánh giá hỗ trợ','user_rating_comment':'Nhận xét phiên hỗ trợ','is_public':'Cho phép hiển thị công khai','admin_reply':'Nội dung phản hồi quản trị','replied_at':'Thời điểm phản hồi','service_type':'Loại dịch vụ được đánh giá',
'plan_name':'Tên giáo án','day_name':'Tên ngày tập','day_of_week':'Ngày trong tuần','week_number':'Tuần thực hiện','current_week':'Tuần hiện tại','duration_weeks':'Thời lượng giáo án (tuần)','estimated_weeks':'Số tuần ước tính','week_start_date':'Ngày bắt đầu tuần','order_index':'Thứ tự bài tập','sets':'Số hiệp','reps':'Số lần lặp','duration_seconds':'Thời lượng (giây)','completion_percent':'Tỷ lệ hoàn thành bài','completion_rate':'Tỷ lệ hoàn thành buổi','sets_completed':'Số hiệp đã hoàn thành','reps_completed':'Số lần lặp đã hoàn thành','weight_used_kg':'Mức tạ đã sử dụng (kg)','is_completed':'Đã hoàn thành','logged_at':'Thời điểm ghi log','session_date':'Ngày tập','scheduled_time':'Giờ tập dự kiến','check_in_time':'Thời điểm bắt đầu tập','check_out_time':'Thời điểm kết thúc tập','duration_minutes':'Thời lượng (phút)','total_calories_burned':'Tổng năng lượng tiêu hao','checkout_body_fat':'Tỷ lệ mỡ lúc kết thúc','checkout_weight':'Cân nặng lúc kết thúc','is_last_session_of_week':'Buổi cuối tuần','is_custom':'Buổi tự tạo','custom_session_name':'Tên buổi tự tạo',
'current_mana':'Sức bền hiện tại','max_mana':'Sức bền tối đa','last_mana_regen_date':'Ngày hồi sức bền gần nhất','last_training_date':'Ngày tập gần nhất','required_max_session_mana_cost':'Chi phí sức bền lớn nhất cần cho một buổi','fitness_score':'Điểm thể lực','body_type':'Phân loại thể trạng','target_level':'Trình độ mục tiêu','target_metric_type':'Chỉ tiêu mục tiêu dạng số: 0 chống đẩy, 1 plank, 2 squat','target_baseline_value':'Giá trị ban đầu của chỉ tiêu','target_current_value':'Giá trị hiện tại của chỉ tiêu','target_goal_value':'Giá trị mục tiêu','target_achieved':'Đã đạt mục tiêu','starting_bmi':'BMI khi bắt đầu','starting_weight':'Cân nặng khi bắt đầu','confirmed_schedule_dows':'Các ngày trong tuần đã xác nhận','is_template':'Là giáo án mẫu','is_ai_generated':'Cờ giáo án được sinh tự động','is_fitness_improvement':'Giáo án cải thiện thể lực','weight_adjustment_note':'Ghi chú điều chỉnh tạ','weight_updated_week':'Tuần cập nhật tạ','base_weight_kg':'Mức tạ gốc (kg)','current_weight_kg':'Mức tạ hiện tại (kg)','recommended_weight_kg':'Mức tạ đề xuất (kg)','current_recommended_weight_kg':'Mức tạ đề xuất hiện tại (kg)',
'membership_type':'Loại hội viên','payment_status':'Trạng thái thanh toán','start_date':'Ngày bắt đầu','end_date':'Ngày kết thúc','invoice_type':'Loại hóa đơn','cosmetic_item_code':'Mã vật phẩm trang trí','cosmetic_code':'Mã vật phẩm sở hữu','purchased_at':'Thời điểm mua','regenerate_count':'Số lần tạo lại thanh toán','momo_order_id':'Mã đơn MoMo','momo_request_id':'Mã yêu cầu MoMo','pay_url':'URL thanh toán','deeplink':'Liên kết mở ứng dụng','result_message':'Thông báo kết quả','stage':'Giai đoạn thể trạng nhân vật','aura_tier':'Cấp hiệu ứng hào quang','current_streak':'Chuỗi duy trì hiện tại','missed_streak':'Chuỗi bỏ lỡ','web_count':'Số mạng nhện hiển thị','last_calculated_at':'Thời điểm tính trạng thái gần nhất','equipped_hair':'Tóc đang trang bị','equipped_shirt':'Áo đang trang bị','equipped_pants':'Quần đang trang bị',
'shift_date':'Ngày làm việc','planned_start':'Giờ bắt đầu dự kiến','planned_end':'Giờ kết thúc dự kiến','check_in_at':'Thời điểm vào ca','check_out_at':'Thời điểm kết ca','cash_at_start':'Tiền mặt đầu ca','cash_counted':'Tiền mặt kiểm đếm','expected_cash':'Tiền mặt dự kiến','cash_difference':'Chênh lệch tiền mặt','pos_bank_revenue':'Doanh thu chuyển khoản tại quầy','pos_cash_revenue':'Doanh thu tiền mặt tại quầy','handover_note':'Ghi chú bàn giao'}
for pre,vn in [('attachment','Tệp đính kèm'),('reply_attachment','Tệp đính kèm phản hồi')]:
    for suf,lab in [('size','kích thước'),('name','tên'),('type','loại'),('url','đường dẫn')]:labels[pre+'_'+suf]=vn+': '+lab
for pref,lab in [('default','mặc định')]:
    for suf,vn in [('duration_seconds','Thời lượng (giây)'),('reps','Số lần lặp'),('sets','Số hiệp')]:labels[pref+'_'+suf]=vn+' '+lab
for c,v in [('endurance_score','sức bền'),('flexibility_score','độ linh hoạt'),('maintenance_score','duy trì'),('muscle_gain_score','tăng cơ'),('weight_loss_score','giảm cân')]:labels[c]='Điểm phù hợp mục tiêu '+v
labels['calories_burned']='Năng lượng tiêu hao của bài tập'
for c,v in [('difficulty_adjustment','độ khó'),('sets_adjustment','số hiệp'),('reps_adjustment','số lần lặp'),('exercises_adjustment','số bài tập')]:labels[c]='Mức điều chỉnh '+v
missing=sorted({c['name'] for t in tables.values() for c in t['columns']}-set(labels))
assert not missing,missing

def constraints(t,c):
    a=['PK'] if c['pk'] else []
    if c['identity']:a+=['Tự tăng']
    a+=['NULL' if c['nullable'] else 'NOT NULL']
    if c['unique']:a+=['UNIQUE']
    for child,col,parent in fks:
        if (child,col)==(t,c['name']):a+=['FK → '+parent+'.'+tables[parent]['pk'][0]]
    for child,col,parent in logical:
        if (child,col)==(t,c['name']):a+=['Logic → '+parent+' (không FK)']
    if 'default ' in c['extra']:a+=['DEFAULT '+re.search(r'default (\w+)',c['extra'])[1]]
    if 'check ' in c['extra']:a+=['CHECK 0..2']
    return '; '.join(a)

corrections=[
('Trang 14, kế hoạch','Đổi mốc “17 entity” thành 41 entity JPA, sinh 43 bảng vật lý gồm 2 bảng liên kết/collection.'),
('Trang 21–22, mục 2.3.1','Bổ sung đầy đủ thực thể hiện có. Hai bảng nối được mô tả rõ ở mô hình logic/vật lý; không coi là hai lớp @Entity riêng.'),
('Trang 22–29, mục 2.3.2','Thay/bổ sung các quan hệ theo danh sách FK và liên kết logic trong tài liệu này. User–Food không có bảng nối hay FK hiện tại; thao tác tham khảo không phải quan hệ N:N được lưu.'),
('Trang 27, ChatMessage','ChatMessage là hội thoại USER/BOT theo user_id. Hỗ trợ người dùng–quản trị dùng support_sessions và support_messages.'),
('Trang 29, EnduranceTest','user_id có UNIQUE và NOT NULL: một người dùng có tối đa một bản ghi kiểm tra hiện tại, không phải bảng lịch sử nhiều lần kiểm tra.'),
('Trang 30 và 33, mục 2.4/2.6','Vẽ lại ERD. Sơ đồ vật lý chứa đủ 43 bảng; dùng các sơ đồ nhóm để đọc được các cột. Chỉ các liên kết FK thật nằm trong ERD vật lý.'),
('Trang 31–32, mục 2.5','Bổ sung nhân viên, cửa hàng, ca làm, nhân vật, cấu hình, hỗ trợ. Mô tả vai trò theo ADMIN/STAFF/USER; không tự thêm vai trò huấn luyện viên nếu chưa triển khai.'),
('Trang 34–35, mục 3.1','Thay danh sách bảng và hình cơ sở dữ liệu bằng đủ 43 bảng, sử dụng tên snake_case đúng database.'),
('Trang 111–122, Phụ lục B','Thay toàn bộ đặc tả cũ bằng 43 đặc tả trong phần 4. Ví dụ roles dùng role_name, không có name/description; PK kiểu BIGINT thay cho INT ở các bảng đang dùng Long.'),
('Các phần yêu cầu, Use Case, UI, triển khai, kiểm thử, hướng dẫn','Cập nhật các chức năng cửa hàng đã làm, chọn phân loại ở cả /shop và /app/shop, quản lý ca nhân viên, địa chỉ, đơn COD, báo cáo. Không ghi ví điện tử đã kiểm thử giao dịch thực tế.')]

md=['# Thông số cập nhật báo cáo: đủ 43 bảng','',
'Đối chiếu mã nguồn GYM_Management và DDL Hibernate của kiểm thử H2 tạm. Ngày lập: 17/09/2026. Tài liệu dùng để thay/bổ sung báo cáo PDF SD-41_DATNSP26 (1).pdf; chưa sửa trực tiếp PDF hay database đang dùng.',
'','## 1. Số lượng và quy ước','',
'Có 43 bảng vật lý = 41 lớp @Entity + 2 bảng nối/collection. Đợt nâng cấp cửa hàng tăng từ 39 lên 43, không phải từ 17 lên 43; số 17 là phần tài liệu cũ chưa đầy đủ. Bốn bảng thêm: customer_addresses, customer_product_states, shop_order_events, shop_inventory_movements.',
'','Kiểu dữ liệu trong phụ lục là DDL H2 thực tế của kiểm thử: Java Long thường là BIGINT; Double là FLOAT(53); ngày giờ là DATE/TIME(6)/TIMESTAMP(6). Không tự đổi thành DECIMAL/INT để làm đẹp tài liệu. NULL chỉ khả năng để trống tại database, không thay thế các kiểm tra đầu vào của dịch vụ. Giá trị khởi tạo trong Java không đồng nghĩa SQL DEFAULT.',
'','Có 39 ràng buộc FK vật lý. Các ID liên kết bằng dịch vụ được liệt kê riêng, không vẽ thành FK đã tồn tại. Hai bảng phụ: variant_attribute_values có PK ghép (value_id, variant_id); voucher_scope_products không có PK/UNIQUE trong DDL hiện tại. Không tự thêm ràng buộc trong báo cáo.',
'','## 2. Vị trí cần sửa trong PDF','']
for loc,txt in corrections:md+=['### '+loc,txt,'']
md+=['## 3. Danh mục toàn bộ bảng','']
for group,names in groups.items():
    md+=['### '+group,'','| Bảng | Chức năng | Số cột |','|---|---|---|']
    md += [f'| {n} | {purposes[n]} | {len(tables[n]["columns"])} |' for n in names]
    md+=['']
md+=['## 4. Đặc tả toàn bộ 43 bảng','']
for idx,t in enumerate([n for ns in groups.values() for n in ns],1):
    data=tables[t]
    md+= [f'### 4.{idx}. {t}',purposes[t]+'.','Nguồn mã: '+sources[t]+'.','']
    if data['unique']:md+=['Ràng buộc duy nhất: '+ '; '.join('('+', '.join(u)+')' for u in data['unique'])+'.','']
    if not data['pk']:md+=['Bảng này chưa có khóa chính trong DDL hiện tại.','']
    md+=['| Cột | Kiểu H2 | Ý nghĩa | Ràng buộc |','|---|---|---|---|']
    for c in data['columns']:
        typ='ENUM' if c['type'].startswith('enum ') else c['type'].upper()
        md += ['| '+' | '.join([c['name'],typ,labels[c['name']],constraints(t,c)])+' |']
    md+=['']
    for c in data['columns']:
        if c['type'].startswith('enum '):md+=['Miền giá trị '+c['name']+': '+', '.join(re.findall(r"'([^']+)'",c['type']))+'.','']

md+=['## 5. Thông số ERD','',
'Trong quan hệ dưới đây, bảng cha có 0..N dòng con nếu cột FK không duy nhất; có 0..1 dòng con nếu FK đồng thời UNIQUE hoặc là toàn bộ PK. Mỗi dòng con tham chiếu 1 cha nếu FK NOT NULL, hoặc 0..1 cha nếu FK cho phép NULL. Đây là bội số của cấu trúc vật lý, không tự suy diễn bắt buộc theo nghiệp vụ.','',
'### 5.1. Các khóa ngoại thực tế','', '| Bảng con.cột | Bảng cha.cột | Số cha trên mỗi con | Số con trên mỗi cha |','|---|---|---|---|']
mermaid=['erDiagram']
for t,d in tables.items():
    mermaid+=['    '+t+' {']
    for c in d['columns']:
        typ='enum' if c['type'].startswith('enum ') else c['type'].split('(')[0]
        keys=[]
        if c['pk']:keys+=['PK']
        if any(a==t and b==c['name'] for a,b,_ in fks):keys+=['FK']
        if c['unique']:keys+=['UK']
        mermaid+=['        '+typ+' '+c['name']+(' '+','.join(keys) if keys else '')]
    mermaid+=['    }']
for t,c,p in fks:
    cd=next(x for x in tables[t]['columns'] if x['name']==c)
    one=([c] in tables[t]['unique'] or [c]==tables[t]['pk'])
    md+=[f'| {t}.{c} | {p}.{tables[p]["pk"][0]} | '+('0..1' if cd['nullable'] else '1')+' | '+('0..1' if one else '0..N')+' |']
    mermaid+=['    '+p+(' |o' if cd['nullable'] else ' ||')+'--'+('o|' if one else 'o{')+' '+t+' : "'+c+'"']
md+=['','### 5.2. Liên kết bằng ID trong ứng dụng, chưa có FK','', '| Bảng.cột | Đối tượng liên quan |','|---|---|']
md +=[f'| {t}.{c} | {p} |' for t,c,p in logical]
md+=['','notifications.ref_id là tham chiếu đa loại, phụ thuộc ref_type; không thể gán chung một FK. Các mã dạng chuỗi như voucher_code hoặc cosmetic_code cần chú thích theo nghiệp vụ, không tự vẽ quan hệ FK.','',
'ERD_43_bang.mmd chứa đủ 43 bảng và toàn bộ cột cùng 39 FK. schema_43_bang.sql là DDL tham chiếu từ H2 tạm, không phải migration để chạy lên dữ liệu đang dùng. Để bố trí bản in: giữ một hình tổng quan đủ bảng, sau đó phóng to theo 6 nhóm tại phần 3; các bảng tham chiếu chéo có thể lặp trên hình nhóm và ghi “tham chiếu”.','',
'## 6. Thông số cập nhật các phần chức năng','',
'### Yêu cầu và Use Case',
'Tác nhân: khách vãng lai, hội viên (USER), nhân viên (STAFF), quản trị (ADMIN); cổng thanh toán là hệ thống bên ngoài. Bổ sung các ca sử dụng: tìm/lọc sản phẩm; xem chi tiết/chọn biến thể; quản lý giỏ; yêu thích/đã xem/so sánh; quản lý địa chỉ; đặt đơn và chọn thanh toán; theo dõi/hủy đơn hợp lệ; đánh giá có ảnh; quản lý thuộc tính/biến thể; xử lý đơn; bán tại quầy; xem báo cáo và nhật ký kho. Các ca đang có được cập nhật, không tạo bản sao chỉ vì thay giao diện.',
'','### Giao diện và tuyến đường',
'/shop: danh mục công khai, giỏ khách vãng lai, phân loại. /shop/product/:id: chi tiết, bộ ảnh, biến thể, đánh giá và sản phẩm liên quan. /app/shop: cửa hàng hội viên, địa chỉ nhận, đặt đơn, lịch sử và đánh giá. /admin/shop: sản phẩm và xử lý đơn. /admin/product-attributes: thuộc tính, giá trị và biến thể. /admin/shop-reports: thống kê, tồn thấp và nhật ký. /staff/orders và /tra-cuu-don: thông tin đơn và phân loại.',
'','### Đơn hàng và thanh toán',
'COD: CONFIRMED → PREPARING → SHIPPING → DELIVERED → COMPLETED; ghi paid_at khi nhân viên xác nhận giao/thu tiền. Thanh toán chờ có PENDING_PAYMENT; QR hết hạn sau 15 phút. Khách chỉ hủy khi đơn thỏa điều kiện dịch vụ. Trừ tồn khi đặt đơn, hoàn tồn khi hủy/hết hạn đúng một lần. Nhật ký lưu thay đổi từ thời điểm triển khai, không có lịch sử giả cho đơn cũ.',
'',
'MoMo/ZaloPay cần cấu hình merchant và callback HTTPS. Callback kiểm tra chữ ký, thông tin thương nhân và số tiền; xử lý lặp có tính idempotent. Trang khách quay về không tự xác nhận đã trả tiền. Thanh toán đến sau hủy/hết hạn được ghi để đối soát, không tự hoàn tiền. Chưa kiểm thử giao dịch thực tế với merchant.',
'','### Tối ưu bảng và quy tắc lưu trữ',
'Một bảng customer_product_states dùng chung yêu thích và lần xem gần nhất; UNIQUE(user_id, product_id). Ảnh sản phẩm/đánh giá là JSON được converter lưu trong VARCHAR(20000), không có bảng ảnh hay kiểu JSON gốc. Đơn lưu ảnh chụp thông tin sản phẩm và biến thể để giữ lịch sử. Biến thể ngừng bán bằng cờ active. Báo cáo tính từ dữ liệu đơn, không thêm bảng tổng hợp.',
'','### Báo cáo doanh thu',
'Lọc ngày theo created_at của đơn. Chỉ cộng đơn có paid_at, loại CANCELLED/EXPIRED. Doanh thu gồm phí vận chuyển và đã trừ giảm giá; bảng sản phẩm bán chạy tính tiền dòng trước giảm giá/phí vận chuyển. Tồn thấp là ≤ 5. Chưa có tính lợi nhuận theo giá vốn.',
'','### Kiểm thử và phạm vi đã xác minh',
'Kết quả có sẵn của đợt nâng cấp: ShopUpgradeTest 9/9, ShopPaymentTest 4/4; frontend build thành công. Toàn bộ backend 32/33, còn một lỗi kiểm thử bộ sinh giáo án được ghi nhận trước đó; không viết “toàn bộ kiểm thử đều đạt”. Số 43 được kiểm chứng bằng schema H2 tạm; không có thao tác migration trên database dữ liệu thật trong lần lập tài liệu này.',
'','### Cách đưa vào báo cáo',
'Thay Phụ lục B bằng phần 4 đầy đủ; thay danh sách bảng mục 3.1.2 bằng phần 3; dựng lại các hình ERD từ phần 5 và tệp Mermaid. Sau khi chèn, cập nhật mục lục, số hình/bảng và các tham chiếu trang trong Word. Số trang ở phần 2 là số trang PDF gốc, sẽ thay đổi khi thêm phụ lục. Giữ nội dung thuật toán tập luyện trừ các mô tả tên trường/quan hệ sai với schema này.']
(OUT/'Thong_so_43_bang.md').write_text('\n'.join(md),encoding='utf-8')
(OUT/'ERD_43_bang.mmd').write_text('\n'.join(mermaid),encoding='utf-8')
(OUT/'schema_43_bang.json').write_text(json.dumps(dict(tables=tables,foreign_keys=fks,logical_links=logical,sources=sources),ensure_ascii=False,indent=2),encoding='utf-8')
(OUT/'schema_43_bang.sql').write_text('-- DDL tham chieu H2 tu kiem thu; KHONG phai migration.\n'+'\n'.join(re.sub(r'^Hibernate: ','',x)+';' for x in raw.splitlines()),encoding='utf-8')

doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.65)
for style in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
    st=doc.styles[style];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(10)
doc.styles['Normal'].paragraph_format.space_after=Pt(5)
doc.styles['Title'].font.size=Pt(23)
doc.styles['Heading 1'].font.size=Pt(16);doc.styles['Heading 2'].font.size=Pt(12)
footer=sec.footer.paragraphs[0];footer.alignment=2
footer.add_run('GYM_Management | Đặc tả 43 bảng | ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
def add_table(lines):
    rows=[[v.strip() for v in l.strip('|').split('|')] for l in lines if not l.startswith('|---')]
    table=doc.add_table(rows=1,cols=len(rows[0]));table.autofit=False
    widths=([1.75,1.05,1.6,2.55] if len(rows[0])==4 else [2.5,3.8,.65])
    if len(rows[0])==2:widths=[3.6,3.35]
    if rows[0][0]=='Bảng con.cột':widths=[2.4,2.4,1.05,1.1]
    for i,w in enumerate(widths):table.columns[i].width=Inches(w)
    for j,vals in enumerate(rows):
        row=table.rows[0] if j==0 else table.add_row()
        for i,value in enumerate(vals):
            cell=row.cells[i];cell.width=Inches(widths[i]);cell.text=value
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
                for r in p.runs:r.font.size=Pt(9);r.bold=j==0
        trpr=row._tr.get_or_add_trPr();keep=OxmlElement('w:cantSplit');trpr.append(keep)
        if j==0:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
            for cell in row.cells:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EEEEEE');cell._tc.get_or_add_tcPr().append(shade)
    props=table._tbl.tblPr;borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'CCCCCC');borders.append(b)
    props.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for edge in ['top','left','bottom','right']:
        e=OxmlElement('w:'+edge);e.set(qn('w:w'),'65');e.set(qn('w:type'),'dxa');margins.append(e)
    props.append(margins)
i=0
while i<len(md):
    line=md[i]
    if line.startswith('|'):
        buf=[]
        while i<len(md) and md[i].startswith('|'):buf.append(md[i]);i+=1
        add_table(buf);continue
    if line.startswith('# '):doc.add_paragraph(line[2:],'Title')
    elif line.startswith('## '):
        p=doc.add_heading(line[3:],1)
        if line.startswith('## 4.') or line.startswith('## 5.'):p.paragraph_format.page_break_before=True
    elif line.startswith('### '):doc.add_heading(line[4:],2)
    elif line:doc.add_paragraph(line)
    i+=1
doc.save(OUT/'Thong_so_cap_nhat_43_bang.docx')
print(json.dumps({'tables':len(tables),'columns':sum(len(x['columns']) for x in tables.values()),'foreign_keys':len(fks),'logical_links':len(logical),'output':str(OUT)},ensure_ascii=False))
