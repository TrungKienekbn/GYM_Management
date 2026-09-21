from docx import Document
import re
p=r'C:\Users\anh15\Downloads\SD-41_DATNSP26.docx'; d=Document(p)
print('paras',len(d.paragraphs),'tables',len(d.tables),'images',len(d.inline_shapes))
for term in ['GYM PRO','Gymify','Khách vãng lai','STAFF','Staff','ROLE_STAFF','Cổng thanh toán','EnduranceTest','UserProfile','PetProfile','19','Mua hàng trực tuyến','Quản trị và vận hành','Phân công và thực hiện']:
 hits=[]
 for i,p0 in enumerate(d.paragraphs):
  if term.lower() in p0.text.lower(): hits.append((i,' '.join(p0.text.split())[:300]))
 print('\n',term,len(hits))
 for h in hits[:6]: print(h)
print('\nTABLE ACTORS/UC')
for ti,t in enumerate(d.tables):
 txt=' | '.join(' '.join(c.text.split()) for row in t.rows for c in row.cells)
 if any(k in txt.lower() for k in ['khách vãng lai','cổng thanh toán','use case','phân công và thực hiện','mua hàng trực tuyến']): print(ti,txt[:1800])
