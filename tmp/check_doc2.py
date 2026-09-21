from docx import Document
import re, zipfile
p=r'C:\Users\anh15\Downloads\SD-41_DATNSP26_Hoan_Thanh (1).docx'
d=Document(p)
print('paras',len(d.paragraphs),'tables',len(d.tables),'sections',len(d.sections),'images',len(d.inline_shapes))
terms=['GYM PRO','Gymify','19 use case','Use case tổng quát','Khách vãng lai','Staff','Nhân viên','UserProfile','EnduranceTest','PetProfile','Phân công','Bán hàng tại quầy']
for term in terms:
 hits=[]
 for i,p0 in enumerate(d.paragraphs):
  if term.lower() in p0.text.lower(): hits.append((i,' '.join(p0.text.split())[:300]))
 print('\nTERM',term,'COUNT',len(hits))
 for h in hits[:8]: print(h)
# extract relationship section relevant paragraphs
print('\nRELATIONSHIP CHECK')
for i in range(370,545):
 t=' '.join(d.paragraphs[i].text.split())
 if any(x.lower() in t.lower() for x in ['User – UserProfile','User – EnduranceTest','User – PetProfile','một User có','UserProfile','EnduranceTest','PetProfile']): print(i,t[:600])
print('\nAPPENDIX HEADINGS')
for i in range(1130,1410):
 t=' '.join(d.paragraphs[i].text.split())
 if re.match(r'7\.\d+',t): print(i,t[:250])
